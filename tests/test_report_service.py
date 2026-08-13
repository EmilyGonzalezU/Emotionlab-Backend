import unittest
from copy import deepcopy
from io import BytesIO

from openpyxl import load_workbook

from services.report_service import ReportService
from services.session_service import STATUS_COMPLETED, SessionService


class FakeFirebase:
    def __init__(self):
        self.documents = {}

    def server_timestamp(self):
        return "SERVER_TIMESTAMP"

    def create_document(self, collection, document_id, data):
        self.documents.setdefault(collection, {})[document_id] = deepcopy(data)

    def get_document(self, collection, document_id):
        data = self.documents.get(collection, {}).get(document_id)
        return deepcopy(data) if data else None

    def list_documents(self, collection):
        return [deepcopy(data) for data in self.documents.get(collection, {}).values()]


class ReportServiceTests(unittest.TestCase):
    def setUp(self):
        sessions = SessionService(FakeFirebase())
        sessions.create_session(
            "session-1",
            {
                "session_data": {
                    "fechaInicio": "2026-07-17 10:00:00",
                    "fechaFin": "2026-07-17 10:05:00",
                    "difficulty_level": "dificil",
                    "respuestas": [{"idPregunta": "edad", "respuestaLabel": "16"}],
                    "eventos": [{"nombre": "session_started", "fase": "inicio"}],
                }
            },
            status=STATUS_COMPLETED,
            prediction={"bienestar": "SI", "probabilidad": 0.93, "confianza": 0.97},
            fecha_inicio="2026-07-17 10:00:00",
        )
        self.reports = ReportService(sessions)

    def test_generates_pdf_and_excel(self):
        pdf = self.reports.generate_individual_pdf("session-1")
        excel = self.reports.generate_sessions_excel()

        self.assertTrue(pdf.startswith(b"%PDF"))
        worksheet = load_workbook(BytesIO(excel))["Sesiones"]
        self.assertEqual(worksheet.max_row, 2)
        self.assertEqual(worksheet["A2"].value, "session-1")
        self.assertEqual(worksheet["I2"].value, "SI")
